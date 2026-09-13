"""
Step 15：輸出一份 Word 報告，說明「版本未經核對」的那 44 張卡片。

背景：卡圖覆蓋率已經是 15,846／15,846（100%），沒有任何一張卡是「載不出圖」
的。但其中 44 張的圖片來自 tcgcollector 補圖包，而這 44 張在 TCGdex 與來源
網站之間編號規則不同，只核對得到「卡包 + 卡名」，核對不到「印刷版本」。
其中 10 張在來源甚至有一個以上的版本，系統依規則挑了一張。

這份 Word 就是給人工覆核用的：每張都附縮圖、卡包、卡號、語言、來源連結，
有其他版本的會把所有版本並列，方便直接比對後回 App 裡換掉。
"""
import io
import json
import os
from datetime import datetime

from PIL import Image
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.dirname(HERE)
PROJECT_DIR = os.path.dirname(DATA_DIR)
CANDIDATES_PATH = os.path.join(DATA_DIR, "image_candidates.json")
DECISIONS_PATH = os.path.join(HERE, "logs", "candidate_decisions.json")
CARDS_PATH = os.path.join(DATA_DIR, "cards.json")
SPECIES_PATH = os.path.join(DATA_DIR, "species.json")
SETS_PATH = os.path.join(DATA_DIR, "sets.json")
OUT_PATH = os.path.join(PROJECT_DIR, "PTCG圖鑑_待人工覆核卡片清單.docx")

THUMB_WIDTH_PX = 320
THUMB_DOC_CM = 2.6

LANG_LABEL = {"en": "英文版", "ja": "日文版", "zh-Hant": "繁體中文版", "zh-Hans": "簡體中文版"}


def thumb(path_rel):
    """把卡圖縮成 320px 寬的 JPEG 放進記憶體，避免整份 Word 變成幾十 MB。"""
    abs_path = os.path.join(PROJECT_DIR, path_rel.replace("/", os.sep))
    if not os.path.exists(abs_path):
        return None
    with Image.open(abs_path) as im:
        im = im.convert("RGB")
        ratio = THUMB_WIDTH_PX / im.width
        im = im.resize((THUMB_WIDTH_PX, max(1, int(im.height * ratio))), Image.LANCZOS)
        buf = io.BytesIO()
        # 用 JPEG 而不是 PNG：卡圖是照片類內容，PNG 一張就 250KB，
        # 58 張會讓整份 Word 變成十幾 MB，寄送與開啟都不方便。
        im.save(buf, format="JPEG", quality=82, optimize=True)
    buf.seek(0)
    return buf


def main():
    candidates = json.load(open(CANDIDATES_PATH, encoding="utf-8"))
    decisions = {d["cardId"]: d for d in json.load(open(DECISIONS_PATH, encoding="utf-8"))}
    cards = {c["id"]: c for c in json.load(open(CARDS_PATH, encoding="utf-8"))}
    species = {s["id"]: s for s in json.load(open(SPECIES_PATH, encoding="utf-8"))}
    sets_meta = json.load(open(SETS_PATH, encoding="utf-8"))

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft JhengHei"
    style.font.size = Pt(10)

    doc.add_heading("PTCG 圖鑑 — 待人工覆核卡片清單", level=0)
    p = doc.add_paragraph()
    p.add_run(f"產生時間：{datetime.now().strftime('%Y-%m-%d %H:%M')}").italic = True

    doc.add_heading("這份文件在講什麼", level=1)
    doc.add_paragraph(
        "目前 App 內 15,846 張卡片全部都有圖片可以顯示，覆蓋率 100%，"
        "沒有任何一張卡是「圖載不出來」的。"
    )
    doc.add_paragraph(
        "但下列 44 張卡片的圖片來自 tcgcollector 補圖包，而這幾張在 TCGdex 與來源網站之間"
        "編號規則不同（或同一個產品裡就有多種印刷版本），只核對得到「卡包」與「卡名」，"
        "核對不到「印刷版本」。圖已經依指示直接套用，這份清單是留給你人工覆核用的。",
    )
    warn = doc.add_paragraph()
    r = warn.add_run(
        f"其中 {sum(1 for d in decisions.values() if d['alternatesNotUsed'])} 張在來源有一個以上的版本，"
        "系統挑了一張（規則：取來源編號最小者；編號相同時取清單第一張）。"
        "這幾張在下面會把所有版本並列，如果挑錯了，到該卡的詳細頁按「換版本」再按「改用這張」就能換。"
    )
    r.bold = True
    r.font.color.rgb = RGBColor(0xB0, 0x50, 0x00)

    doc.add_heading("怎麼換成別的版本", level=1)
    for i, line in enumerate([
        "開啟 App（雙擊「啟動圖鑑.bat」），進到該張卡所屬寶可夢的詳細頁。",
        "找到那張卡，卡包資訊那行右邊會有一顆小按鈕「換版本 N」（N 是還有幾個其他版本）。平常候選圖是收合的，不會佔版面。",
        "按下「換版本」展開，比對這份文件裡的縮圖與來源連結，確認要哪一版後按「改用這張」。",
        "想換回原本的，按卡片圖片上的「×（移除自訂圖片）」即可恢復成預設套用的版本。",
    ], 1):
        doc.add_paragraph(f"{i}. {line}", style="List Number" if False else None)

    # ---- 有多個版本的，優先列出 ----
    multi = [cid for cid, d in decisions.items() if d["alternatesNotUsed"]]
    single = [cid for cid, d in decisions.items() if not d["alternatesNotUsed"]]

    def card_header(cid):
        card = cards.get(cid, {})
        set_info = sets_meta.get(f"{card.get('language')}:{card.get('setId')}", {})
        dex = card.get("dexNumbers") or []
        zh = "、".join(species[d]["nameZh"] for d in dex if d in species)
        title = card.get("name") or cid
        if zh:
            title = f"{zh}（{title}）"
        return title, set_info.get("setName") or card.get("setId"), card

    def add_card_section(cid, with_alternates):
        d = decisions[cid]
        info = candidates.get(cid, {})
        title, set_name, card = card_header(cid)
        doc.add_heading(title, level=2)

        meta = doc.add_paragraph()
        meta.add_run("卡包：").bold = True
        meta.add_run(f"{set_name}　")
        meta.add_run("卡號：").bold = True
        meta.add_run(f"{card.get('cardNumber') or '—'}　")
        meta.add_run("語言：").bold = True
        meta.add_run(f"{LANG_LABEL.get(card.get('language'), card.get('language'))}　")
        meta.add_run("卡片 ID：").bold = True
        meta.add_run(cid)

        note = doc.add_paragraph()
        nr = note.add_run(info.get("note") or "")
        nr.font.size = Pt(9)
        nr.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

        cands = info.get("candidates") or []
        applied_path = info.get("appliedPath")
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        hdr = table.rows[0].cells
        for c, text in zip(hdr, ["圖片", "狀態", "來源資訊"]):
            c.text = ""
            run = c.paragraphs[0].add_run(text)
            run.bold = True

        for cand in cands:
            row = table.add_row().cells
            buf = thumb(cand["path"])
            cell_p = row[0].paragraphs[0]
            cell_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if buf:
                cell_p.add_run().add_picture(buf, width=Cm(THUMB_DOC_CM))
            else:
                cell_p.add_run("（縮圖產生失敗）")

            is_applied = cand["path"] == applied_path
            st = row[1].paragraphs[0].add_run("目前使用中" if is_applied else "未使用")
            st.bold = is_applied
            if is_applied:
                st.font.color.rgb = RGBColor(0x1B, 0x6E, 0x1B)

            info_p = row[2].paragraphs[0]
            info_p.add_run(f"{cand.get('title') or ''}\n").bold = True
            info_p.add_run(f"來源卡包：{cand.get('setName') or '—'}\n")
            info_p.add_run(f"來源編號：{cand.get('number') or '—'}\n")
            info_p.add_run(f"來源頁：{cand.get('sourceUrl') or '—'}")
            for r_ in info_p.runs:
                r_.font.size = Pt(8)

        if with_alternates:
            why = doc.add_paragraph()
            wr = why.add_run(f"挑選理由：{d['pickReason']}")
            wr.font.size = Pt(9)
            wr.italic = True

    doc.add_page_break()
    doc.add_heading(f"一、來源有多個版本，需要你確認（{len(multi)} 張）", level=1)
    doc.add_paragraph("這些卡片同一個產品裡就有不只一種印刷（例如牌組限定的藍框版與一般版，"
                      "或四副牌各自的牌組符號）。請比對縮圖後決定要用哪一版。")
    for cid in multi:
        add_card_section(cid, True)

    doc.add_page_break()
    doc.add_heading(f"二、來源只有一個版本，已直接套用（{len(single)} 張）", level=1)
    doc.add_paragraph("這些卡片來源只找得到一張圖，已直接採用。卡包與卡名核對過，"
                      "但編號規則無法對應，所以仍列在這裡供覆核。")
    for cid in single:
        add_card_section(cid, False)

    # 使用者可能正開著這份 Word（Windows 會鎖住檔案），這時不要讓整個流程掛掉，
    # 也不要硬搶檔案，改存成「(更新版)」讓使用者自己決定要不要取代。
    out_path = OUT_PATH
    try:
        doc.save(out_path)
    except PermissionError:
        base, ext = os.path.splitext(OUT_PATH)
        out_path = f"{base}(更新版){ext}"
        doc.save(out_path)

    print(json.dumps({
        "output": os.path.relpath(out_path, PROJECT_DIR),
        "savedToFallbackName": out_path != OUT_PATH,
        "multiVersionCards": len(multi),
        "singleVersionCards": len(single),
        "totalCards": len(decisions),
        "fileSizeKB": round(os.path.getsize(out_path) / 1024),
    }, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
